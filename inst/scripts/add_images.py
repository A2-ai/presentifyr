import os
import argparse
import json
from pptx import Presentation
from PIL import Image
from py_logger import get_logger
from pptx_utils import (
    USABLE_PLACEHOLDER_TYPES,
    PRESERVE_PLACEHOLDER_INDICES,
    find_footer_placeholder,
    load_metadata_for_image,
    format_slide_notes,
)

## 914400 EMU = 1 inch
EMU_PER_INCH = 914400

## DPI used by PowerPoint auto-scaling (matches R add_images.R:165-166)
DEFAULT_DPI = 300


def get_image_dimensions_inches(path, dpi=DEFAULT_DPI):
    """Read pixel dimensions via Pillow and convert to inches at given DPI.

    Mirrors magick::image_read + image_info in add_images.R:154-169.
    """
    with Image.open(path) as img:
        width_px, height_px = img.size
    return width_px / dpi, height_px / dpi


def find_layout_by_name(prs, name):
    """Find a slide layout by name with normalization.

    Mirrors layout matching in add_images.R:42-55.
    Normalizes by lowering case, replacing underscores with spaces, and stripping whitespace.
    """
    logger = get_logger()

    if not name or not isinstance(name, str):
        ## Default to index 1 (second layout, matching officer default)
        layout = prs.slide_layouts[1]
        logger.debug(f"No layout name specified, using default layout: {layout.name}")
        return layout

    normalized_name = name.strip().lower().replace('_', ' ')
    logger.debug(f"Looking for layout matching: '{normalized_name}'")

    for layout in prs.slide_layouts:
        normalized_layout = layout.name.strip().lower().replace('_', ' ')
        if normalized_layout == normalized_name:
            logger.debug(f"Matched layout: '{layout.name}'")
            return layout

    raise ValueError(f"Invalid layout name '{name}' not found in template.")


def get_usable_placeholders(layout):
    """Get usable placeholders (types 7/18) sorted by position (top, left).

    Mirrors add_images.R:61-78.
    Returns list of dicts with idx, type, left, top, width, height (all in EMU).
    """
    placeholders = []
    for shape in layout.placeholders:
        ph_type = shape.placeholder_format.type
        if ph_type in USABLE_PLACEHOLDER_TYPES:
            placeholders.append({
                'idx': shape.placeholder_format.idx,
                'type': int(ph_type),
                'left': shape.left,
                'top': shape.top,
                'width': shape.width,
                'height': shape.height,
            })

    ## Sort by position: top-to-bottom, then left-to-right
    placeholders.sort(key=lambda p: (p['top'], p['left']))
    return placeholders


def clone_missing_placeholders(slide, layout):
    """Clone all layout placeholders missing from the slide.

    python-pptx's add_slide() only clones a subset of placeholders (title,
    body, content, picture) to the new slide. Other placeholders like footer
    (idx=11), date (idx=10), and slide number (idx=12) are inherited visually
    from the layout/master but have no actual XML element on the slide.
    PowerPoint's own "Insert Slide" clones all of them. This function does
    the same so the slide matches what a user would get from the UI.
    """
    from copy import deepcopy
    logger = get_logger()

    existing_indices = {
        shape.placeholder_format.idx for shape in slide.placeholders
    }

    for shape in layout.placeholders:
        idx = shape.placeholder_format.idx
        if idx not in existing_indices:
            sp_clone = deepcopy(shape._element)
            slide.shapes._spTree.append(sp_clone)
            logger.debug(f"Cloned placeholder idx={idx} ({shape.name}) from layout to slide")


def clear_non_image_placeholders(slide, usable_indices):
    """Clear text in non-image placeholders, preserving footer and slide number.

    Mirrors add_images.R:116-126.  The R code uses officer::ph_with(value = "")
    which blanks the placeholder text but keeps the element in the slide XML.
    Removing the element entirely would strip template formatting/decorations.
    """
    logger = get_logger()
    for shape in list(slide.placeholders):
        idx = shape.placeholder_format.idx
        if idx not in usable_indices and idx not in PRESERVE_PLACEHOLDER_INDICES:
            if shape.has_text_frame:
                shape.text_frame.clear()
            logger.debug(f"Cleared placeholder idx={idx}")


def place_image_in_placeholder(slide, image_path, ph_info, alt_text):
    """Insert image into placeholder, scale to fit, and center within bounds.

    Uses placeholder.insert_picture() so the placeholder is properly consumed,
    then adjusts dimensions for scale-to-fit + center (matching the R/officer
    behavior from add_images.R:154-191).
    """
    from pptx.shapes.placeholder import PicturePlaceholder

    logger = get_logger()

    ph_idx = ph_info['idx']

    ## Access the actual placeholder on the slide
    ph = slide.placeholders[ph_idx]

    ## Store original placeholder bounds before insertion
    ph_left = ph.left
    ph_top = ph.top
    ph_width = ph.width
    ph_height = ph.height

    ## Content placeholders (type 7) are generic SlidePlaceholder objects
    ## that lack insert_picture(). Cast to PicturePlaceholder to get the method.
    ## Type 18 (Picture) placeholders already have it natively.
    if not hasattr(ph, 'insert_picture'):
        ph = PicturePlaceholder(ph._element, ph._parent)

    ## Insert picture INTO the placeholder (properly consumes it)
    pic = ph.insert_picture(image_path)

    ## Reset any auto-cropping applied by insert_picture's fill behavior
    pic.crop_top = 0
    pic.crop_left = 0
    pic.crop_bottom = 0
    pic.crop_right = 0

    ## Calculate scale-to-fit using image aspect ratio
    with Image.open(image_path) as img:
        img_w_px, img_h_px = img.size

    img_aspect = img_w_px / img_h_px
    ph_aspect = ph_width / ph_height

    if img_aspect > ph_aspect:
        ## Image is wider relative to placeholder — constrained by width
        final_w = ph_width
        final_h = int(round(ph_width / img_aspect))
    else:
        ## Image is taller relative to placeholder — constrained by height
        final_h = ph_height
        final_w = int(round(ph_height * img_aspect))

    ## Center within original placeholder bounds
    pic.left = ph_left + (ph_width - final_w) // 2
    pic.top = ph_top + (ph_height - final_h) // 2
    pic.width = final_w
    pic.height = final_h

    ## Set alt-text using {prfy}:<key> pattern
    pic._element.nvPicPr.cNvPr.set("descr", alt_text)

    logger.debug(f"Placed image in placeholder idx={ph_idx}, "
                 f"size: {final_w / EMU_PER_INCH:.2f} x {final_h / EMU_PER_INCH:.2f} inches")


def set_slide_notes_plain(slide, text):
    """Set plain text slide notes.

    Mirrors add_images.R:261.
    """
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = text


def set_slide_notes_styled(slide, text, font_settings):
    """Set styled slide notes using python-pptx font API.

    Mirrors add_images.R:239-242.
    """
    from pptx.util import Pt
    from pptx.dml.color import RGBColor

    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame

    ## Clear existing content
    text_frame.clear()

    lines = text.split('\n')
    for i, line in enumerate(lines):
        if i == 0:
            para = text_frame.paragraphs[0]
        else:
            para = text_frame.add_paragraph()

        run = para.add_run()
        run.text = line

        _apply_font_settings(run.font, font_settings)

        ## Bold headers
        if line.startswith("## "):
            run.font.bold = True


def set_footer_text_plain(shape, text):
    """Insert plain text into footer placeholder.

    Mirrors add_images.R:253-258.
    """
    shape.text_frame.text = text


def set_footer_text_styled(shape, text, font_settings):
    """Insert styled text into footer placeholder.

    Mirrors add_images.R:232-236.
    """
    from pptx.util import Pt
    from pptx.dml.color import RGBColor

    text_frame = shape.text_frame
    text_frame.clear()

    lines = text.split('\n')
    for i, line in enumerate(lines):
        if i == 0:
            para = text_frame.paragraphs[0]
        else:
            para = text_frame.add_paragraph()

        run = para.add_run()
        run.text = line

        _apply_font_settings(run.font, font_settings)

        ## Bold headers
        if line.startswith("## "):
            run.font.bold = True


def _apply_font_settings(font, font_settings):
    """Apply font settings dict to a python-pptx font object."""
    from pptx.util import Pt
    from pptx.dml.color import RGBColor

    font.size = Pt(font_settings.get('font_size', 8))
    font.name = font_settings.get('font_name', 'Calibri')
    font.bold = font_settings.get('bold', False)
    font.italic = font_settings.get('italic', False)
    font.underline = font_settings.get('underline', False)

    color_hex = font_settings.get('font_color', '#000000')
    if color_hex.startswith('#'):
        color_hex = color_hex[1:]
    font.color.rgb = RGBColor.from_string(color_hex)


def add_images(output_pptx, config, base_pptx=None):
    """Main orchestration function for adding images to a PPTX.

    Mirrors add_images.R:24-270.

    Args:
        output_pptx: Path to write the output PPTX file
        config: dict with slide_layout_name, slide_groups, slide_positions,
                font_settings, image_keys
        base_pptx: Optional path to a base PPTX template
    """
    logger = get_logger()
    logger.debug("Starting add_images Python function")

    ## Open or create presentation
    if base_pptx and os.path.exists(base_pptx):
        logger.info(f"Using base PowerPoint template: {base_pptx}")
        prs = Presentation(base_pptx)
    else:
        logger.warning("No valid PowerPoint template found. Creating a blank presentation.")
        prs = Presentation()

    ## Find the target layout
    slide_layout_name = config.get('slide_layout_name')
    layout = find_layout_by_name(prs, slide_layout_name)
    logger.debug(f"Selected layout: {layout.name}")

    ## Get usable placeholders
    usable_phs = get_usable_placeholders(layout)
    if len(usable_phs) == 0:
        raise ValueError("No usable placeholder found (Content Placeholder or Picture Placeholder)")

    logger.debug(f"Found {len(usable_phs)} usable placeholder(s)")

    ## Detect footer placeholder availability from layout
    has_footer = any(
        shape.placeholder_format.idx == 11
        for shape in layout.placeholders
        if shape.is_placeholder
    )
    if has_footer:
        logger.debug("Footer placeholder detected in layout")

    ## Get usable placeholder indices for clearing
    usable_indices = {ph['idx'] for ph in usable_phs}

    slide_groups = config['slide_groups']
    slide_positions = config['slide_positions']
    font_settings = config.get('font_settings', {})
    image_keys = config.get('image_keys', {})
    abbrev_defs = config.get('abbreviation_definitions', None)
    meta_type_definitions = config.get('meta_type_definitions', None)

    logger.info(f"Total slides to create: {len(slide_groups)}")

    for slide_idx, (slide_files, slide_pos) in enumerate(zip(slide_groups, slide_positions), start=1):
        ## auto_unbox in R serializes single-element vectors as scalars;
        ## normalize back to lists so zip() works uniformly
        if isinstance(slide_files, str):
            slide_files = [slide_files]
        if isinstance(slide_pos, (int, float)):
            slide_pos = [slide_pos]
        logger.debug(f"Creating slide {slide_idx} of {len(slide_groups)} with {len(slide_files)} image(s)")

        slide = prs.slides.add_slide(layout)

        ## Clone any layout placeholders that add_slide() didn't copy,
        ## so the slide matches what PowerPoint's own "Insert Slide" produces
        clone_missing_placeholders(slide, layout)

        ## Clear non-image placeholders (title, etc.) but keep usable + footer + slidenum
        clear_non_image_placeholders(slide, usable_indices)

        ## Collect metadata from all images on this slide
        slide_metadata_list = {}

        ## Place each image in corresponding placeholder
        for img_idx, (file_path, target_pos) in enumerate(zip(slide_files, slide_pos)):
            ## Clamp to valid range (1-indexed from R, convert to 0-indexed)
            ph_row_idx = min(max(target_pos - 1, 0), len(usable_phs) - 1)
            ph_info = usable_phs[ph_row_idx]

            logger.debug(f"Placing image {img_idx + 1} in placeholder idx={ph_info['idx']}")

            ## Build alt-text using {prfy}:<key> pattern
            alt_text_key = image_keys.get(file_path, os.path.basename(file_path))
            alt_text = f"{{prfy}}:{alt_text_key}"

            place_image_in_placeholder(slide, file_path, ph_info, alt_text)

            ## Collect metadata for this image
            metadata = load_metadata_for_image(file_path)
            if metadata:
                slide_metadata_list[os.path.basename(file_path)] = metadata

        ## Combine metadata from all images into footnote content
        if slide_metadata_list:
            use_styled = len(font_settings) > 0

            ## Build combined notes text
            combined_parts = []
            for name, meta in slide_metadata_list.items():
                combined_parts.append(
                f"## {name}\n"
                f"{format_slide_notes(meta, abbrev_defs, meta_type_definitions)}"
            )
            combined_notes = "\n\n".join(combined_parts)

            footer_shape = find_footer_placeholder(slide)

            if use_styled:
                if footer_shape is not None:
                    set_footer_text_styled(footer_shape, combined_notes, font_settings)
                    logger.debug(f"Added styled footnotes to footer for slide {slide_idx}")
                else:
                    ## Pre-initialize notes to avoid inheriting notes master formatting
                    set_slide_notes_plain(slide, "")
                    set_slide_notes_styled(slide, combined_notes, font_settings)
                    logger.debug(f"Added styled combined notes for slide {slide_idx}")
            else:
                if footer_shape is not None:
                    set_footer_text_plain(footer_shape, combined_notes)
                    logger.debug(f"Added footnotes to footer for slide {slide_idx}")
                else:
                    set_slide_notes_plain(slide, combined_notes)
                    logger.debug(f"Added combined notes for slide {slide_idx}")

    prs.save(output_pptx)
    logger.info(f"PowerPoint saved as {output_pptx}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Add images to PowerPoint slides using python-pptx"
    )
    parser.add_argument('-o', '--output_pptx', type=str, required=True,
                        help="Output pptx file path")
    parser.add_argument('-c', '--config', type=str, required=True,
                        help="Path to JSON config file")
    parser.add_argument('-b', '--base_pptx', type=str, default=None,
                        help="Base PowerPoint template file path")

    args = parser.parse_args()

    with open(args.config, 'r') as f:
        config = json.load(f)

    add_images(
        output_pptx=args.output_pptx,
        config=config,
        base_pptx=args.base_pptx
    )
