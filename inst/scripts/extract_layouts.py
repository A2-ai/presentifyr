import os
import argparse
import re
from pptx import Presentation
from PIL import Image, ImageDraw
from py_logger import get_logger

## Placeholder types that can accept images
## 7 = OBJECT (Content Placeholder) - universal content placeholder
## 18 = PICTURE (Picture Placeholder) - image-specific placeholder
USABLE_PLACEHOLDER_TYPES = {7, 18}

def generate_layout_images(base_pptx, output_dir):
    logger = get_logger()
    logger.debug("Starting generate layout images Python function")

    prs = Presentation(base_pptx)

    max_width_px = 800  ## Hardcoded based on testing

    slide_width_emu = prs.slide_width
    slide_height_emu = prs.slide_height
    logger.debug(f"Slide master dimensions (EMUs): width={slide_width_emu}, height={slide_height_emu}")

    width_px = max_width_px
    aspect_ratio = slide_height_emu / float(slide_width_emu)
    height_px = int(width_px * aspect_ratio)

    layout_count = 0
    for i, layout in enumerate(prs.slide_layouts):
        layout_name = layout.name
        safe_layout_name = re.sub(r'[^\w\-_]', '_', layout_name)  # Replace spaces & special characters

        usable_placeholder_count = sum(
            1 for shape in layout.shapes
            if shape.is_placeholder and shape.placeholder_format.type in USABLE_PLACEHOLDER_TYPES
        )

        logger.debug(f"Layout '{layout_name}' (Index {i}) has {usable_placeholder_count} usable placeholder(s) (types {USABLE_PLACEHOLDER_TYPES})")

        if usable_placeholder_count < 1:
            logger.debug(f"Skipping layout '{layout_name}' (Index {i}) since it has no usable placeholders")
            continue

        logger.info(f"Processing layout '{layout_name}' (Index {i}) ({usable_placeholder_count} usable placeholder(s))")

        img = Image.new("RGB", (width_px, height_px), "white")
        draw = ImageDraw.Draw(img)

        draw.text((10, 10), layout_name, fill="black")

        x_scale = width_px / float(slide_width_emu)
        y_scale = height_px / float(slide_height_emu)

        for shape in layout.shapes:
            if shape.is_placeholder:
                left_px = int(shape.left * x_scale)
                top_px = int(shape.top * y_scale)
                right_px = int((shape.left + shape.width) * x_scale)
                bottom_px = int((shape.top + shape.height) * y_scale)

                fill_color = None

                if shape.placeholder_format.type in USABLE_PLACEHOLDER_TYPES:
                    fill_color = (255, 200, 200)  ## Soft red for usable placeholders

                draw.rectangle(
                    [(left_px, top_px), (right_px, bottom_px)],
                    outline="black",
                    fill=fill_color,
                    width=2
                )

        img_path = os.path.join(output_dir, f"{safe_layout_name}.png")
        logger.debug(f"Saving thumbnail as '{img_path}'")
        img.save(img_path)
        layout_count += 1

    logger.info(f"Generated {layout_count} layout image(s) in '{output_dir}'")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Extract and visualize PowerPoint slide layouts that have at least one usable placeholder (types 7 or 18)")
    parser.add_argument('-b', '--base_pptx', type=str, required=True, help="Base PowerPoint template")
    parser.add_argument('-o', '--output_dir', type=str, required=True, help="Directory to save layout images")
    
    args = parser.parse_args()

    generate_layout_images(base_pptx=args.base_pptx, output_dir=args.output_dir)
