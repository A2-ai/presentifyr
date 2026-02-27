import os
import re
import argparse
import json
import hashlib
from pptx import Presentation
from py_logger import get_logger
from pptx_utils import find_footer_placeholder, load_metadata_for_image, format_slide_notes


def compute_file_hash(file_path):
    """Compute MD5 hash of a file on disk."""
    hash_md5 = hashlib.md5()
    with open(file_path, 'rb') as f:
        for chunk in iter(lambda: f.read(8192), b''):
            hash_md5.update(chunk)
    return hash_md5.hexdigest()


def compute_blob_hash(blob):
    """Compute MD5 hash of binary data (image blob from PPTX)."""
    return hashlib.md5(blob).hexdigest()


def compute_text_hash(text):
    """Compute MD5 hash of text content."""
    return hashlib.md5(text.encode('utf-8')).hexdigest()


def should_replace_image(shape, new_image_path, logger):
    """Determine if image should be replaced based on hash comparison."""
    try:
        existing_hash = compute_blob_hash(shape.image.blob)
        new_hash = compute_file_hash(new_image_path)

        if existing_hash == new_hash:
            return False, "unchanged (hashes match)"
        else:
            return True, "content changed"
    except Exception as e:
        logger.warning(f"Could not compute hash: {e}. Will replace image.")
        return True, f"hash computation failed ({e})"


def normalize_text_for_comparison(text):
    """Normalize text for hash comparison.

    Strips trailing whitespace from each line and trailing blank lines,
    so officer-created styled content (which may have extra empty paragraphs)
    matches the Python-generated text.
    """
    lines = text.split('\n')
    # Strip trailing whitespace per line
    lines = [line.rstrip() for line in lines]
    # Strip trailing blank lines
    while lines and lines[-1] == '':
        lines.pop()
    return '\n'.join(lines)


def should_update_footer(footer_shape, new_text, logger):
    """Determine if footer text should be updated based on hash comparison."""
    try:
        existing_text = normalize_text_for_comparison(footer_shape.text or "")
        new_normalized = normalize_text_for_comparison(new_text)

        existing_hash = compute_text_hash(existing_text)
        new_hash = compute_text_hash(new_normalized)

        if existing_hash == new_hash:
            return False, "unchanged (hashes match)"
        else:
            return True, "content changed"
    except Exception as e:
        logger.warning(f"Could not compare footer: {e}. Will update.")
        return True, f"comparison failed ({e})"


def should_update_notes(slide, new_notes_text, logger):
    """Determine if slide notes should be updated based on hash comparison."""
    try:
        notes_slide = slide.notes_slide
        text_frame = notes_slide.notes_text_frame
        if text_frame is None:
            return True, "no existing notes"

        existing_text = normalize_text_for_comparison(text_frame.text or "")
        new_normalized = normalize_text_for_comparison(new_notes_text)

        existing_hash = compute_text_hash(existing_text)
        new_hash = compute_text_hash(new_normalized)

        if existing_hash == new_hash:
            return False, "unchanged (hashes match)"
        else:
            return True, "content changed"
    except Exception as e:
        logger.warning(f"Could not compare notes: {e}. Will update.")
        return True, f"comparison failed ({e})"


def classify_line(text):
    """Classify a line for block-aware paragraph matching.

    Returns:
        tuple: (kind, key)
            - ("blank", None) for empty lines
            - ("header", <full header text>) for lines starting with ## 
            - ("kv", <normalized key>) for Key: value lines
            - ("text", <normalized text>) for non-empty, non-header, non-key lines
    """
    stripped = (text or "").strip()

    if stripped == "":
        return "blank", None
    if stripped.startswith("## "):
        return "header", stripped
    if ":" in stripped:
        key = stripped.split(":", 1)[0].strip().casefold()
        if key != "":
            return "kv", key
    return "text", stripped.casefold()


def _pick_unmatched(candidates, used_ids):
    """Return the first unmatched paragraph from a candidate list."""
    for para in candidates:
        if id(para) not in used_ids:
            return para
    return None


def update_text_preserve_formatting(text_frame, new_text):
    """Update text frame content while preserving existing run formatting.

    Matches old paragraphs to new lines within each image block.
    - Headers match by exact "## ..." text
    - Metadata lines match by normalized key before the first colon
    - Other lines match by normalized full text

    Matched paragraphs get text updated with formatting preserved.
    Missing lines are inserted unstyled.
    """
    from lxml import etree
    logger = get_logger()

    ns = '{http://schemas.openxmlformats.org/drawingml/2006/main}'

    new_lines = new_text.split('\n')
    while new_lines and new_lines[-1].strip() == '':
        new_lines.pop()

    existing_paras = list(text_frame.paragraphs)

    # --- Group existing paragraphs into blocks keyed by header text ---
    # Each block stores:
    #   - header paragraph
    #   - paragraphs indexed by normalized key (list to handle duplicates)
    #   - paragraphs indexed by normalized full text (list fallback)
    # Also track blank separator paragraphs between blocks
    old_blocks = {}       # header_text -> block data
    old_separators = {}   # header_text -> [blank paras before this block]
    current_header = None
    current_blanks = []

    for p in existing_paras:
        text = (p.text or '').rstrip()
        kind, value = classify_line(text)

        if kind == 'header':
            # Stash any accumulated blanks as separators for this block
            if current_header is not None:
                old_separators[value] = current_blanks
            current_blanks = []
            current_header = value
            old_blocks[current_header] = {
                'header': p,
                'by_key': {},
                'by_text': {}
            }
        elif kind == 'blank':
            current_blanks.append(p)
        elif current_header is not None:
            block = old_blocks[current_header]
            if kind == 'kv':
                block['by_key'].setdefault(value, []).append(p)
                if len(block['by_key'][value]) == 2:
                    logger.warning(
                        "Duplicate key '%s' found in block '%s'; using first unmatched paragraph.",
                        value,
                        current_header
                    )
            else:
                block['by_text'].setdefault(value, []).append(p)

    # --- Group new lines into blocks by header ---
    new_blocks = []  # list of (header_text, [lines])
    current_new = []
    current_new_header = None
    for line in new_lines:
        kind, value = classify_line(line)
        if kind == 'header':
            if current_new_header is not None:
                new_blocks.append((current_new_header, current_new))
            current_new = []
            current_new_header = value
        if kind != 'blank':
            current_new.append(line)
    if current_new_header is not None:
        new_blocks.append((current_new_header, current_new))

    # --- Detach all existing paragraphs ---
    txBody = text_frame._txBody
    for p_elem in txBody.findall(f'{ns}p'):
        txBody.remove(p_elem)

    # --- Rebuild paragraphs in order ---
    for block_i, (header_text, block_lines) in enumerate(new_blocks):
        old_block = old_blocks.get(header_text, {})
        used_ids = set()

        # Add blank separator between blocks
        if block_i > 0:
            sep_paras = old_separators.get(header_text, [])
            if sep_paras:
                for sp in sep_paras:
                    txBody.append(sp._p)
            else:
                # No old separator — create a plain blank paragraph
                p_elem = etree.SubElement(txBody, f'{ns}p')

        for new_line in block_lines:
            kind, value = classify_line(new_line)
            old_para = None

            if kind == 'header':
                candidate = old_block.get('header')
                if candidate is not None and id(candidate) not in used_ids:
                    old_para = candidate
            elif kind == 'kv':
                candidates = old_block.get('by_key', {}).get(value, [])
                old_para = _pick_unmatched(candidates, used_ids)
            elif kind == 'text':
                candidates = old_block.get('by_text', {}).get(value, [])
                old_para = _pick_unmatched(candidates, used_ids)

            if old_para is not None:
                # Reuse existing paragraph — update text, keep formatting
                p_elem = old_para._p
                if old_para.runs:
                    old_para.runs[0].text = new_line
                    for run in old_para.runs[1:]:
                        p_elem.remove(run._r)
                else:
                    r_elem = etree.SubElement(p_elem, f'{ns}r')
                    t_elem = etree.SubElement(r_elem, f'{ns}t')
                    t_elem.text = new_line
                txBody.append(p_elem)
                used_ids.add(id(old_para))
            else:
                # Missing (user deleted it) — insert plain, unstyled
                p_elem = etree.SubElement(txBody, f'{ns}p')
                r_elem = etree.SubElement(p_elem, f'{ns}r')
                t_elem = etree.SubElement(r_elem, f'{ns}t')
                t_elem.text = new_line


def sync_images(input_pptx, output_pptx, image_dict):
    logger = get_logger()
    logger.debug(f"Starting sync images Python function")

    presentation = Presentation(input_pptx)
    logger.debug(f"Using input .pptx file: {input_pptx}")

    start_pattern = r'\{prfy\}\:'
    end_pattern = r'\.[^.]+$'
    magic_pattern = re.compile(start_pattern + '.*?' + end_pattern)

    # Track statistics for summary
    stats = {
        'images_replaced': 0,
        'images_skipped': 0,
        'images_not_found': 0,
        'notes_updated': 0,
        'notes_skipped': 0,
        'footer_updated': 0,
        'footer_skipped': 0
    }

    logger.info("Scanning slides for image and footnote sync")

    for slide_index, slide in enumerate(presentation.slides, start=1):
        matched_images = []  # All images with {prfy}: marker and valid paths
        replacements = []    # Images that need replacement (hash differs)

        for shape in slide.shapes:

            # Extract alt_text based on shape type
            if shape.shape_type == 14:  # GROUP
                alt_text = shape._element._nvXxPr.cNvPr.attrib.get("descr", "")
            elif shape.shape_type == 13:  # PICTURE
                alt_text = shape._element.nvPicPr.cNvPr.attrib.get("descr", "")
            else:
                continue

            logger.debug(f"Slide {slide_index}: Checking image alt-text: {alt_text}")

            match = magic_pattern.match(alt_text)
            if match:
                figure_name = match.group(0).split(':')[1]
                # Try exact match first (for relative paths), then basename match
                image_path = image_dict.get(figure_name)
                if not image_path:
                    # Fall back to matching by basename (for backward compatibility)
                    for key, path in image_dict.items():
                        if os.path.basename(key) == figure_name or os.path.basename(path) == figure_name:
                            image_path = path
                            break

                if image_path and os.path.exists(image_path):
                    # Track all matched images for footnote processing
                    matched_images.append((shape, image_path, alt_text))

                    # Check if image content has changed using hash comparison
                    should_replace, reason = should_replace_image(shape, image_path, logger)

                    if should_replace:
                        logger.info(f"Slide {slide_index}: Queuing replacement for {figure_name} ({reason})")
                        replacements.append((shape, image_path, alt_text))
                    else:
                        logger.debug(f"Slide {slide_index}: Skipping figure {figure_name} - {reason}")
                        stats['images_skipped'] += 1
                else:
                    logger.warning(f"Slide {slide_index}: No matching image found for {figure_name} or file does not exist")
                    stats['images_not_found'] += 1

        # Process figure replacements
        for shape, image_path, alt_text in replacements:
            left = shape.left
            top = shape.top
            width = shape.width
            height = shape.height

            slide.shapes._spTree.remove(shape._element)
            new_pic = slide.shapes.add_picture(image_path, left, top, width, height)
            new_pic._element.nvPicPr.cNvPr.set("descr", alt_text)

            logger.debug(f"Slide {slide_index}: Replaced image from {image_path}")
            stats['images_replaced'] += 1

        # Process footnotes independently (using ALL matched images, not just replacements)
        if matched_images:
            # Collect metadata from all matched images on this slide
            slide_metadata_list = {}
            for _, image_path, _ in matched_images:
                metadata = load_metadata_for_image(image_path)
                if metadata:
                    slide_metadata_list[os.path.basename(image_path)] = metadata

            if slide_metadata_list:
                # Combine notes from all images (matching add_images.R behavior)
                combined_notes = "\n\n".join(
                    f"## {name}\n{format_slide_notes(meta)}"
                    for name, meta in slide_metadata_list.items()
                )

                # Check for footer placeholder first; fall back to slide notes
                footer_shape = find_footer_placeholder(slide)

                if footer_shape is not None:
                    should_update, reason = should_update_footer(footer_shape, combined_notes, logger)
                    if should_update:
                        try:
                            update_text_preserve_formatting(footer_shape.text_frame, combined_notes)
                            logger.info(f"Slide {slide_index}: Updated footer ({reason})")
                            stats['footer_updated'] += 1
                        except Exception as e:
                            logger.warning(f"Slide {slide_index}: Could not update footer - {e}")
                    else:
                        logger.debug(f"Slide {slide_index}: Skipping footer - {reason}")
                        stats['footer_skipped'] += 1
                else:
                    # No footer placeholder — fall back to slide notes
                    should_update, reason = should_update_notes(slide, combined_notes, logger)
                    if should_update:
                        try:
                            notes_slide = slide.notes_slide
                            text_frame = notes_slide.notes_text_frame
                            if text_frame is not None:
                                update_text_preserve_formatting(text_frame, combined_notes)
                                logger.info(f"Slide {slide_index}: Updated notes ({reason})")
                                stats['notes_updated'] += 1
                            else:
                                logger.warning(f"Slide {slide_index}: No notes text frame available")
                        except Exception as e:
                            logger.warning(f"Slide {slide_index}: Could not update notes - {e}")
                    else:
                        logger.debug(f"Slide {slide_index}: Skipping notes - {reason}")
                        stats['notes_skipped'] += 1

        if not matched_images:
            logger.warning(f"Slide {slide_index}: No matching alt-text found")

    logger.info(f"Sync complete: {stats['images_replaced']} images replaced, "
                f"{stats['images_skipped']} images unchanged, "
                f"{stats['footer_updated']} footers updated, "
                f"{stats['footer_skipped']} footers unchanged, "
                f"{stats['notes_updated']} notes updated, "
                f"{stats['notes_skipped']} notes unchanged")

    presentation.save(output_pptx)
    logger.debug(f"PowerPoint saved as {output_pptx}")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Sync images within input pptx file by monkey-patching placeholders.")
    parser.add_argument('-i', '--input_pptx', type=str, required=True, help="Input pptx file path")
    parser.add_argument('-o', '--output_pptx', type=str, required=True, help="Output pptx file path")
    parser.add_argument('-d', '--image_dict', type=str, required=True, help="Path to JSON file containing image dictionary")

    args = parser.parse_args()

    with open(args.image_dict, 'r') as f:
        image_dict = json.load(f)

    sync_images(
        input_pptx=args.input_pptx,
        output_pptx=args.output_pptx,
        image_dict=image_dict
    )
